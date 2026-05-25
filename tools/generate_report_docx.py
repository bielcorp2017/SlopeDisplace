"""옹벽 변위 분석 Word 리포트 생성.

이전 PDF 리포트는 경사(tilt) 중심이었고, 이 워드 리포트는
경사 · 수평 · 수직 세 종류의 변위를 모두 다룬다.

페이지/섹션 구성:
    표지     — 현장명, 측정 헤드라인 3종
    1. 측정 개요
    2. 수직봉(plumb) 측정 결과
    3. 경사 분석 (회귀)
    4. 수평 변위 분석 (분포 + 높이 프로파일)
    5. 수직 변위 분석 (분포 + 높이 프로파일)
    6. 종합 결론

사용:
    python tools/generate_report_docx.py <dataset> [--site "현장명"] [--out report.docx]
"""

from __future__ import annotations

import argparse
import io
import json
import sys
from datetime import datetime
from pathlib import Path

import numpy as np
import open3d as o3d
from scipy.spatial import cKDTree

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager

from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# --- Korean font for matplotlib ---
_FONT_PATH = r"C:\Windows\Fonts\malgun.ttf"
if Path(_FONT_PATH).exists():
    font_manager.fontManager.addfont(_FONT_PATH)
    plt.rcParams["font.family"] = "Malgun Gothic"
plt.rcParams["axes.unicode_minus"] = False

ROOT = Path(__file__).resolve().parent.parent


# ============================================================
#  Style helpers
# ============================================================
KOREAN_FONT = "맑은 고딕"
MONO_FONT = "Consolas"
ACCENT = RGBColor(0x1A, 0x4D, 0x80)
TEXT = RGBColor(0x22, 0x22, 0x22)
DIM = RGBColor(0x66, 0x66, 0x66)
RED = RGBColor(0xC0, 0x00, 0x00)
BLUE = RGBColor(0x00, 0x66, 0xAA)
GREEN = RGBColor(0x00, 0x88, 0x00)


def _set_run_font(run, name=KOREAN_FONT, size=10.5, bold=False, color=None):
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_paragraph(doc, text, *, size=10.5, bold=False, color=None,
                  alignment=None, mono=False, space_after=4):
    p = doc.add_paragraph()
    if alignment is not None:
        p.alignment = alignment
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    _set_run_font(run, name=MONO_FONT if mono else KOREAN_FONT,
                  size=size, bold=bold, color=color)
    return p


def add_heading(doc, text, level=1):
    """Custom heading — python-docx default uses Calibri Light w/ no Korean."""
    sizes = {1: 16, 2: 13, 3: 11}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6 if level == 1 else 4)
    run = p.add_run(text)
    _set_run_font(run, size=sizes.get(level, 11),
                  bold=True, color=ACCENT)
    return p


def add_kv_table(doc, items, *, k_width=Cm(5.5), v_width=Cm(10.5)):
    """items: list of (key, value, [color]) tuples."""
    table = doc.add_table(rows=0, cols=2)
    table.autofit = False
    for it in items:
        if len(it) >= 3:
            k, v, c = it[0], it[1], it[2]
        else:
            k, v, c = it[0], it[1], None
        row = table.add_row().cells
        row[0].width = k_width
        row[1].width = v_width
        # key cell
        p = row[0].paragraphs[0]
        run = p.add_run(k)
        _set_run_font(run, size=10, color=DIM)
        # value cell
        p2 = row[1].paragraphs[0]
        run2 = p2.add_run(v)
        _set_run_font(run2, name=MONO_FONT, size=10, bold=False,
                      color=c if c else TEXT)
    return table


def add_image(doc, fig, width_cm=15):
    """matplotlib Figure → PNG → docx 삽입."""
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=140, bbox_inches="tight",
                facecolor="white")
    plt.close(fig)
    buf.seek(0)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run()
    run.add_picture(buf, width=Cm(width_cm))


# ============================================================
#  Stats / charts
# ============================================================
def stats_of(arr):
    """기본 통계 dict."""
    a = np.asarray(arr, dtype=np.float64)
    return {
        "n": int(a.size),
        "min": float(a.min()),
        "max": float(a.max()),
        "mean": float(a.mean()),
        "abs_mean": float(np.abs(a).mean()),
        "median": float(np.median(a)),
        "rms": float(np.sqrt(np.mean(a * a))),
        "p95_abs": float(np.percentile(np.abs(a), 95)),
        "std": float(a.std()),
    }


def fmt_signed(v, d=3):
    if v is None or not np.isfinite(v):
        return "—"
    return f"{v:+.{d}f}"


def histogram_chart(values_mm, label, color, n_bins=80, clip_mm=None):
    fig, ax = plt.subplots(figsize=(7, 3.2))
    v = np.asarray(values_mm, dtype=np.float64)
    if clip_mm:
        v = v[(v >= -clip_mm) & (v <= clip_mm)]
    ax.hist(v, bins=n_bins, color=color, alpha=0.85, edgecolor="white", linewidth=0.3)
    ax.axvline(0, color="#333", lw=0.6, linestyle="--")
    ax.axvline(float(np.mean(v)), color="#c00", lw=1.0, label=f"mean={np.mean(v):+.2f}mm")
    ax.set_xlabel(label)
    ax.set_ylabel("점 개수")
    ax.legend(loc="upper right", fontsize=9)
    ax.grid(alpha=0.3, axis="y")
    fig.tight_layout()
    return fig


def profile_chart(h_m, val_mm, label, color, n_slices=30):
    """높이 구간별 평균±std 프로파일."""
    h = np.asarray(h_m); v = np.asarray(val_mm)
    bins = np.linspace(h.min(), h.max(), n_slices + 1)
    centers = 0.5 * (bins[:-1] + bins[1:])
    means, stds, counts = [], [], []
    for i in range(n_slices):
        m = (h >= bins[i]) & (h < bins[i + 1])
        if m.sum() < 5:
            means.append(np.nan); stds.append(np.nan); counts.append(0); continue
        means.append(np.mean(v[m]))
        stds.append(np.std(v[m]))
        counts.append(int(m.sum()))
    means = np.array(means); stds = np.array(stds)

    fig, ax = plt.subplots(figsize=(7, 3.6))
    valid = ~np.isnan(means)
    ax.fill_betweenx(centers[valid], means[valid] - stds[valid],
                     means[valid] + stds[valid], alpha=0.25, color=color)
    ax.plot(means[valid], centers[valid], "-o", color=color, ms=3, label="평균 ± 1σ")
    ax.axvline(0, color="#333", lw=0.6)
    ax.set_xlabel(label)
    ax.set_ylabel("높이 h (m, plumb 축)")
    ax.legend(loc="best", fontsize=9)
    ax.grid(alpha=0.3)
    fig.tight_layout()
    return fig


def regression_scatter_chart(t):
    fig, ax = plt.subplots(figsize=(7, 3.6))
    h = np.asarray(t["scatter_h_m"]); d = np.asarray(t["scatter_d_mm"])
    ax.scatter(h, d, s=1.2, alpha=0.25, c="#3a7bd5")
    h_lo, h_hi = t["height_range_m"]
    hh = np.linspace(h_lo, h_hi, 100)
    ax.plot(hh, t["alpha_mm_per_m"] * (hh - t["h_ref_m"]) + t["beta_mm"],
            "r-", lw=2, label=f"α={t['alpha_mm_per_m']:+.3f} mm/m  ({t['tilt_change_deg']:+.4f}°)")
    ax.axhline(0, color="#333", lw=0.4)
    ax.set_xlabel("높이 h (m, plumb 축)")
    ax.set_ylabel("signed normal disp (mm) — outward +")
    ax.legend(loc="upper left", fontsize=9)
    ax.grid(alpha=0.3)
    fig.tight_layout()
    return fig


# ============================================================
#  Wall mask + per-point analysis
# ============================================================
def load_target_data(dataset_root: Path, target_stem: str, ref_stem: str,
                     plumb_axis: np.ndarray):
    """target 점 + disp + 옹벽 면 마스크 + plumb 높이 반환."""
    tgt_ply = dataset_root / f"{target_stem}_simple.ply"
    disp_bin = dataset_root / f"{target_stem}_disp.bin"
    ref_ply = dataset_root / f"{ref_stem}_simple.ply"
    pcd_t = o3d.io.read_point_cloud(str(tgt_ply))
    P_t = np.asarray(pcd_t.points, dtype=np.float64)
    disp = np.fromfile(disp_bin, dtype=np.float32).reshape(-1, 4)
    # ref normals
    pcd_r = o3d.io.read_point_cloud(str(ref_ply))
    pcd_r.estimate_normals(o3d.geometry.KDTreeSearchParamHybrid(radius=0.5, max_nn=30))
    pcd_r.orient_normals_consistent_tangent_plane(k=20)
    P_r = np.asarray(pcd_r.points, dtype=np.float64)
    N_r = np.asarray(pcd_r.normals, dtype=np.float64)
    tree = cKDTree(P_r)
    _, idx = tree.query(P_t, k=1, workers=-1)
    n_at_t = N_r[idx]
    dot_plumb = np.abs(n_at_t @ plumb_axis)
    mask_wall = dot_plumb < 0.3
    return P_t, disp, mask_wall, n_at_t


# ============================================================
#  Pages
# ============================================================
def page_cover(doc, site, dataset, tilt, plumb, meta):
    # Title
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(40)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("옹벽 변위 모니터링 분석 리포트")
    _set_run_font(run, size=24, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(site)
    _set_run_font(run, size=18, bold=True, color=ACCENT)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(20)
    run = p.add_run(f"데이터셋  ·  {dataset}")
    _set_run_font(run, size=11, color=DIM)

    # divider
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("━" * 28)
    _set_run_font(run, size=10, color=DIM)

    # Summary
    add_heading(doc, "측정 요약", level=2)

    ds = meta.get("displacement_stats", {}) if meta else {}
    sn = ds.get("signed_normal", {})
    hz = ds.get("horizontal", {})
    vt = ds.get("vertical", {})

    items = [
        ("기준일",          f"{fmt_date(tilt['reference'])}"),
        ("최신일",          f"{fmt_date(tilt['target'])}"),
        ("관측기간",         f"{days_between(tilt['reference'], tilt['target'])} 일"),
        ("plumb 벡터",      f"({plumb['axis'][0]:+.4f}, {plumb['axis'][1]:+.4f}, {plumb['axis'][2]:+.4f})"),
        ("Z축 대비 기울기",   f"{plumb['tilt_from_z_deg']:.4f}°"),
        ("", ""),
    ]
    items.append(("[경사] 추가 기울기 변화",
                  f"{tilt['tilt_change_deg']:+.5f}°  ({tilt['tilt_direction']})",
                  RED if (tilt['tilt_direction']=='OUTWARD' and tilt['significant']) else
                  BLUE if (tilt['tilt_direction']=='INWARD' and tilt['significant']) else GREEN))
    items.append(("[경사] 95% CI",
                  f"[{tilt['tilt_change_ci95_deg'][0]:+.5f}, {tilt['tilt_change_ci95_deg'][1]:+.5f}]°"))
    items.append(("[수평] 평균 |변위|", f"{hz.get('abs_mean',0)*1000:.2f} mm"))
    items.append(("[수평] p95",         f"{hz.get('p95_abs',0)*1000:.2f} mm"))
    items.append(("[수직] 평균 |변위|", f"{vt.get('abs_mean',0)*1000:.2f} mm"))
    items.append(("[수직] p95",         f"{vt.get('p95_abs',0)*1000:.2f} mm"))
    items.append(("", ""))
    items.append(("ICP fitness",       f"{tilt.get('icp_fitness',0)*100:.1f}%"))
    items.append(("정합 신뢰",          "RELIABLE" if tilt.get('registration_reliable') else "UNRELIABLE",
                  GREEN if tilt.get('registration_reliable') else RED))

    add_kv_table(doc, items)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(40)
    run = p.add_run(f"생성일자   {datetime.now():%Y-%m-%d %H:%M}")
    _set_run_font(run, size=9, color=DIM)

    doc.add_page_break()


def page_overview(doc, site, dataset, tilt, plumb):
    add_heading(doc, "1. 측정 개요", level=1)

    add_paragraph(doc,
        f"본 리포트는 '{site}' 현장에 설치된 옹벽의 3차원 점군 스캔 데이터를 "
        f"분석하여, 기준일 대비 옹벽의 변형 양상을 정량적으로 평가한 결과이다.")

    add_paragraph(doc, "분석 대상", bold=True, space_after=2, size=11)
    items = [
        ("대상",       f"{site} 옹벽"),
        ("기준 스캔",   f"{fmt_date(tilt['reference'])}"),
        ("최신 스캔",   f"{fmt_date(tilt['target'])}"),
        ("기간",       f"{days_between(tilt['reference'], tilt['target'])} 일"),
        ("점 개수",    f"{tilt['total_point_count']:,} (옹벽 면 {tilt['wall_face_point_count']:,})"),
    ]
    add_kv_table(doc, items)

    add_paragraph(doc, "분석 항목", bold=True, space_after=2, size=11)
    add_paragraph(doc,
        "  ① 경사(Tilt)     — 옹벽 전체 평균 기울기의 추가 변화",
        size=10)
    add_paragraph(doc,
        "  ② 수평 변위      — 옹벽 facing 방향 ±, 분포 및 높이별 프로파일",
        size=10)
    add_paragraph(doc,
        "  ③ 수직 변위      — Z축(중력) 방향 변위, 침하·융기 양상",
        size=10)

    add_paragraph(doc, "분석 방법론", bold=True, space_after=2, size=11)
    add_paragraph(doc,
        "1) 현장 설치 수직봉을 별도 스캔하여 PCA·슬라이스원피팅·원기둥LSQ "
        "세 방법 교차검증으로 정밀 plumb 벡터를 추출한다.")
    add_paragraph(doc,
        "2) 두 시점의 점군을 FGR + 다중스케일 ICP(1m→1cm) 로 정합한다. "
        "정합 fitness 가 30% 미만이면 정합 신뢰불가로 판정한다.")
    add_paragraph(doc,
        "3) 정합 후 각 점에 대해 변위 벡터를 계산하고, 표면 법선·수평·수직 "
        "성분으로 분해한다(*_disp.bin: signed_normal, magnitude, horizontal, vertical).")
    add_paragraph(doc,
        "4) 옹벽 면 점(|n·plumb|<0.3) 만 마스킹하여, 평탄면(상단·바닥)의 "
        "영향을 배제한다.")
    add_paragraph(doc,
        "5) 경사는 signed_normal 을 plumb 축 높이에 대해 robust 회귀(Huber)하여 "
        "α[mm/m] 와 95% bootstrap CI 를 구한다.")
    add_paragraph(doc,
        "6) 수평·수직은 분포 통계와 높이별 프로파일로 양상을 평가한다.")

    doc.add_page_break()


def page_plumb(doc, plumb):
    add_heading(doc, "2. 수직봉(plumb) 측정 결과", level=1)
    add_paragraph(doc,
        "현장에 설치한 수직봉을 별도 스캔하고, 세 가지 독립 방법으로 축 벡터를 "
        "추출하여 일치 여부를 확인했다. 세 방법이 모두 같은 축을 가리키면 "
        "측정 신뢰의 증거가 된다.")

    val = plumb.get("validation", {})
    items = [
        ("원본 PLY",       Path(plumb["source_ply"]).name),
        ("축 벡터",        f"({plumb['axis'][0]:+.6f}, {plumb['axis'][1]:+.6f}, {plumb['axis'][2]:+.6f})"),
        ("중심",           f"({plumb['center'][0]:.3f}, {plumb['center'][1]:.3f}, {plumb['center'][2]:.3f}) m"),
        ("반경",           f"{plumb['radius_m']*1000:.2f} mm"),
        ("봉 길이",        f"{plumb['rod_length_m']:.3f} m"),
        ("Z축 기울기",     f"{plumb['tilt_from_z_deg']:.4f}°"),
        ("표면 잡음 σ",    f"{plumb['surface_sigma_m']*1000:.2f} mm"),
        ("인라이어 점 수",  f"{plumb['n_inliers']:,}"),
    ]
    add_kv_table(doc, items)

    add_heading(doc, "검증 — 세 방법의 일치도", level=3)
    vitems = [
        ("Cylinder LSQ ↔ PCA",     f"{val.get('cylinder_vs_pca_deg', 0):.4f}°"),
        ("Cylinder LSQ ↔ Slice",   f"{val.get('cylinder_vs_slice_deg', 0):.4f}°"),
        ("PCA ↔ Slice",            f"{val.get('pca_vs_slice_deg', 0):.4f}°"),
        ("선형성 (L)",              f"{val.get('linearity', 0):.4f}"),
        ("planarity (P)",          f"{val.get('planarity', 0):.4f}"),
        ("종합 PROVEN",            "YES" if val.get("overall_proven") else "NO",
         GREEN if val.get("overall_proven") else RED),
    ]
    add_kv_table(doc, vitems)

    add_paragraph(doc,
        "→ 세 방법 모두 1° 이내로 일치, 선형성·평면성 PASS, 표면 잡음이 스캐너 "
        "노이즈 수준이면 측정된 수직 벡터를 신뢰할 수 있다.",
        color=DIM, size=10, space_after=6)

    doc.add_page_break()


def page_tilt(doc, tilt):
    add_heading(doc, "3. 경사(Tilt) 분석", level=1)
    add_paragraph(doc,
        "옹벽 표면점의 signed_normal 변위(=법선 방향 outward 변위)를 plumb 축 "
        "높이 h 에 대해 robust 회귀하면 α[mm/m] 를 얻는다. α 의 atan 이 곧 "
        "기준일 대비 옹벽의 추가 기울기이다.")

    fig = regression_scatter_chart(tilt)
    add_image(doc, fig, width_cm=15)

    add_heading(doc, "회귀 결과", level=3)
    a = tilt["alpha_mm_per_m"]
    deg = tilt["tilt_change_deg"]
    color = (RED if deg > 0 and tilt["significant"] else
             BLUE if deg < 0 and tilt["significant"] else GREEN)
    items = [
        ("α (회귀 기울기)",          f"{a:+.4f} mm/m"),
        ("α 95% CI",               f"[{tilt['alpha_ci95_mm_per_m'][0]:+.4f}, {tilt['alpha_ci95_mm_per_m'][1]:+.4f}]"),
        ("β (h_ref 평균변위)",       f"{tilt['beta_mm']:+.3f} mm"),
        ("h 범위",                  f"{tilt['height_range_m'][0]:.2f} ~ {tilt['height_range_m'][1]:.2f} m"),
        ("잔차 robust σ",            f"{tilt['residual_robust_sigma_mm']:.2f} mm"),
        ("R²",                     f"{tilt['r_squared']:.4f}"),
        ("회귀 사용 옹벽 점 수",       f"{tilt['wall_face_point_count']:,}"),
        ("", ""),
        ("기울기 변화",              f"{deg:+.5f}° ({tilt['tilt_direction']})", color),
        ("95% CI",                  f"[{tilt['tilt_change_ci95_deg'][0]:+.5f}, {tilt['tilt_change_ci95_deg'][1]:+.5f}]°"),
        ("판정",                    tilt["verdict"], color),
    ]
    add_kv_table(doc, items)

    interp = ("옹벽이 outward(앞쪽/하향) 방향으로 미세하게 더 기울고 있음. "
              "변화량은 작지만 95% CI 가 0 을 포함하지 않으므로 통계적으로 유의."
              if tilt["significant"] and a > 0 else
              "옹벽이 inward(뒤쪽/사면쪽) 방향으로 더 기울고 있음. "
              "통계적으로 유의."
              if tilt["significant"] and a < 0 else
              "기울기 변화가 95% CI 안에 0 이 포함되어 통계적으로 유의하지 않음. "
              "잡음 수준의 변화로 해석.")
    add_paragraph(doc, f"해석: {interp}", color=DIM, size=10)

    doc.add_page_break()


def page_horizontal(doc, tilt, P_t, disp, mask_wall, plumb_axis, plumb_origin):
    add_heading(doc, "4. 수평 변위 분석", level=1)
    add_paragraph(doc,
        "수평 변위는 disp[:,2] (옹벽의 XY 노멀 방향 부호 투영, +outward) 를 사용한다. "
        "옹벽이 facing 방향으로 얼마나 이동했는지를 측정한다.")

    horiz_all = disp[:, 2].astype(np.float64) * 1000.0  # mm
    horiz_w = horiz_all[mask_wall]

    s_all = stats_of(horiz_all)
    s_wall = stats_of(horiz_w)

    add_heading(doc, "통계 — 전체 vs 옹벽 면", level=3)
    items = [
        ("전체 점 수",     f"{s_all['n']:,}"),
        ("옹벽 점 수",     f"{s_wall['n']:,}"),
        ("", ""),
        ("[전체] mean",    f"{s_all['mean']:+.3f} mm"),
        ("[전체] |mean|",  f"{s_all['abs_mean']:.3f} mm"),
        ("[전체] σ",       f"{s_all['std']:.3f} mm"),
        ("[전체] p95(|·|)", f"{s_all['p95_abs']:.3f} mm"),
        ("[전체] max",     f"{s_all['max']:+.3f} mm"),
        ("[전체] min",     f"{s_all['min']:+.3f} mm"),
        ("", ""),
        ("[옹벽] mean",    f"{s_wall['mean']:+.3f} mm",
         RED if s_wall['mean'] > 1 else BLUE if s_wall['mean'] < -1 else GREEN),
        ("[옹벽] |mean|",  f"{s_wall['abs_mean']:.3f} mm"),
        ("[옹벽] σ",       f"{s_wall['std']:.3f} mm"),
        ("[옹벽] p95(|·|)", f"{s_wall['p95_abs']:.3f} mm"),
        ("[옹벽] max",     f"{s_wall['max']:+.3f} mm"),
        ("[옹벽] min",     f"{s_wall['min']:+.3f} mm"),
    ]
    add_kv_table(doc, items)

    # histogram (wall only)
    add_heading(doc, "옹벽 면 수평 변위 분포", level=3)
    clip = max(20, abs(s_wall['p95_abs']) * 3)
    fig = histogram_chart(horiz_w, "수평 변위 (mm) — outward +", "#3a7bd5",
                          clip_mm=clip)
    add_image(doc, fig, width_cm=15)

    # height profile
    add_heading(doc, "높이별 수평 변위 프로파일", level=3)
    h_all = (P_t - plumb_origin) @ plumb_axis
    fig = profile_chart(h_all[mask_wall], horiz_w,
                        "옹벽 수평 변위 (mm)", "#3a7bd5")
    add_image(doc, fig, width_cm=15)

    add_paragraph(doc,
        "해석: 평균이 양(+)이면 옹벽이 전반적으로 outward 방향으로, 음(-)이면 "
        "inward 방향으로 이동. 높이 프로파일에서 위쪽 패치만 outward 면 lean. "
        "전 구간 일정한 양수면 평행이동.",
        color=DIM, size=10)

    doc.add_page_break()


def page_vertical(doc, P_t, disp, mask_wall, plumb_axis, plumb_origin):
    add_heading(doc, "5. 수직 변위 분석", level=1)
    add_paragraph(doc,
        "수직 변위는 disp[:,3] (signed_normal 의 Z성분) 을 사용한다. 옹벽 표면의 "
        "Z 방향 움직임 — 침하(-) 또는 융기(+) 양상을 본다.")

    vert_all = disp[:, 3].astype(np.float64) * 1000.0  # mm
    vert_w = vert_all[mask_wall]

    s_all = stats_of(vert_all)
    s_wall = stats_of(vert_w)

    add_heading(doc, "통계 — 전체 vs 옹벽 면", level=3)
    items = [
        ("전체 점 수",     f"{s_all['n']:,}"),
        ("옹벽 점 수",     f"{s_wall['n']:,}"),
        ("", ""),
        ("[전체] mean",    f"{s_all['mean']:+.3f} mm"),
        ("[전체] |mean|",  f"{s_all['abs_mean']:.3f} mm"),
        ("[전체] σ",       f"{s_all['std']:.3f} mm"),
        ("[전체] p95(|·|)", f"{s_all['p95_abs']:.3f} mm"),
        ("[전체] max(융기)", f"{s_all['max']:+.3f} mm"),
        ("[전체] min(침하)", f"{s_all['min']:+.3f} mm"),
        ("", ""),
        ("[옹벽] mean",    f"{s_wall['mean']:+.3f} mm",
         RED if s_wall['mean'] > 1 else BLUE if s_wall['mean'] < -1 else GREEN),
        ("[옹벽] |mean|",  f"{s_wall['abs_mean']:.3f} mm"),
        ("[옹벽] σ",       f"{s_wall['std']:.3f} mm"),
        ("[옹벽] p95(|·|)", f"{s_wall['p95_abs']:.3f} mm"),
        ("[옹벽] max(융기)", f"{s_wall['max']:+.3f} mm"),
        ("[옹벽] min(침하)", f"{s_wall['min']:+.3f} mm"),
    ]
    add_kv_table(doc, items)

    add_heading(doc, "옹벽 면 수직 변위 분포", level=3)
    clip = max(20, abs(s_wall['p95_abs']) * 3)
    fig = histogram_chart(vert_w, "수직 변위 (mm) — 융기 +", "#27ae60",
                          clip_mm=clip)
    add_image(doc, fig, width_cm=15)

    add_heading(doc, "높이별 수직 변위 프로파일", level=3)
    h_all = (P_t - plumb_origin) @ plumb_axis
    fig = profile_chart(h_all[mask_wall], vert_w,
                        "옹벽 수직 변위 (mm)", "#27ae60")
    add_image(doc, fig, width_cm=15)

    add_paragraph(doc,
        "해석: 평균이 양(+)이면 융기, 음(-)이면 침하. 높이별 프로파일이 "
        "균일하면 전체 평행이동, 위쪽에서 큰 음수면 상부 침하·하부 안정 패턴 등.",
        color=DIM, size=10)

    doc.add_page_break()


def page_conclusion(doc, tilt, P_t, disp, mask_wall):
    add_heading(doc, "6. 종합 결론 및 권고", level=1)

    add_heading(doc, "주요 발견", level=3)

    horiz_w = disp[mask_wall, 2].astype(np.float64) * 1000.0
    vert_w = disp[mask_wall, 3].astype(np.float64) * 1000.0
    hm = float(horiz_w.mean()); vm = float(vert_w.mean())
    hpp = float(np.percentile(np.abs(horiz_w), 95))
    vpp = float(np.percentile(np.abs(vert_w), 95))

    a = tilt["alpha_mm_per_m"]
    deg = tilt["tilt_change_deg"]
    days = days_between(tilt['reference'], tilt['target'])

    bullets = []
    sig_tilt = ("OUTWARD" if (tilt["significant"] and a > 0)
                else "INWARD" if (tilt["significant"] and a < 0)
                else "변화 없음")
    bullets.append(f"경사: 기준일 대비 {days}일 사이 {abs(deg):.4f}° {sig_tilt} "
                   f"(95% CI {tilt['tilt_change_ci95_deg'][0]:+.4f}° ~ "
                   f"{tilt['tilt_change_ci95_deg'][1]:+.4f}°)")
    h_dir = "outward" if hm > 0 else "inward"
    bullets.append(f"수평: 옹벽 면 평균 {hm:+.2f} mm ({h_dir}), p95 절대값 {hpp:.2f} mm")
    v_dir = "융기" if vm > 0 else "침하"
    bullets.append(f"수직: 옹벽 면 평균 {vm:+.2f} mm ({v_dir}), p95 절대값 {vpp:.2f} mm")
    bullets.append(f"ICP 정합 fitness {tilt.get('icp_fitness',0)*100:.1f}%, "
                   f"잔차 σ {tilt['residual_robust_sigma_mm']:.2f} mm — "
                   f"{'신뢰 가능' if tilt.get('registration_reliable') else '신뢰 주의'}")

    for b in bullets:
        add_paragraph(doc, "  •  " + b, size=10.5)

    add_heading(doc, "측정 한계", level=3)
    limits = [
        "두 시점 간 *추가* 변위를 측정 — 기준일 자체의 절대 수직도가 아님.",
        "옹벽 디자인 굴곡은 두 스캔에 공통이므로 차분에서 상쇄됨.",
        "ICP 정합 품질이 낮으면 (<30%) 측정값은 옹벽 변화가 아닌 정합 오차일 가능성.",
        "스캐너 잡음 한계(2~5mm) 이하의 변화는 검출 불가.",
    ]
    for l in limits:
        add_paragraph(doc, "  •  " + l, size=10)

    add_heading(doc, "권고 사항", level=3)
    recs = [
        "모니터링 주기: 주 1회 ~ 월 1회 (계측 환경에 따라 조정).",
        f"경고 임계값(예시): 누적 기울기 변화 0.10° 초과, 수평 변위 p95 절대값 "
        f"10 mm 초과, 수직 침하 5 mm 초과 시 현장 점검 권장.",
        "여러 스캔이 누적되면 시계열 추세 분석으로 lean 진행률(°/일) 정량 평가 가능.",
        "ICP 정합 품질이 낮은 회차는 재처리 후 재분석 권장.",
    ]
    for r in recs:
        add_paragraph(doc, "  •  " + r, size=10)


# ============================================================
#  Utility
# ============================================================
def fmt_date(stem: str) -> str:
    if len(stem) == 8 and stem.isdigit():
        return f"{stem[:4]}-{stem[4:6]}-{stem[6:8]}"
    return stem


def days_between(stem_a: str, stem_b: str) -> int:
    try:
        a = datetime.strptime(stem_a, "%Y%m%d")
        b = datetime.strptime(stem_b, "%Y%m%d")
        return (b - a).days
    except Exception:
        return 0


def set_document_defaults(doc):
    """문서 기본 폰트를 맑은 고딕으로."""
    style = doc.styles["Normal"]
    style.font.name = KOREAN_FONT
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), KOREAN_FONT)
    style.font.size = Pt(10.5)
    # 페이지 여백
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)


def add_header_footer(doc, site):
    """헤더에 현장명, 푸터에 페이지 번호."""
    section = doc.sections[0]
    # header
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(site + "  ·  옹벽 변위 분석")
    _set_run_font(run, size=9, color=DIM)
    # footer
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run()
    fldChar = OxmlElement("w:fldChar"); fldChar.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText"); instrText.text = "PAGE"
    fldChar2 = OxmlElement("w:fldChar"); fldChar2.set(qn("w:fldCharType"), "end")
    run._element.append(fldChar); run._element.append(instrText); run._element.append(fldChar2)
    _set_run_font(run, size=9, color=DIM)


# ============================================================
#  main
# ============================================================
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("dataset")
    ap.add_argument("--site", default="현장명 미지정")
    ap.add_argument("--out", default=None)
    ap.add_argument("--data-root", default=str(ROOT / "data"))
    args = ap.parse_args()

    root = Path(args.data_root) / args.dataset
    plumb = json.loads((root / "plumb.json").read_text(encoding="utf-8"))

    tilt_files = sorted(root.glob("wall_tilt_*.json"))
    if not tilt_files:
        print("[ERROR] no wall_tilt_*.json", file=sys.stderr); sys.exit(1)
    tilts = [json.loads(p.read_text(encoding="utf-8")) for p in tilt_files]
    tilts.sort(key=lambda x: x["target"])
    tilt = tilts[-1]  # 최신

    # 최신 스캔의 meta + 점군 + 마스크
    ref_stem = tilt["reference"]
    tgt_stem = tilt["target"]
    meta = json.loads((root / f"{tgt_stem}_meta.json").read_text(encoding="utf-8"))
    plumb_axis = np.asarray(plumb["axis"], dtype=np.float64)
    plumb_axis /= np.linalg.norm(plumb_axis)
    plumb_origin = np.asarray(plumb["center"], dtype=np.float64)

    print(f"[load] {tgt_stem} simple.ply + disp.bin + ref normals ...")
    P_t, disp, mask_wall, _ = load_target_data(root, tgt_stem, ref_stem, plumb_axis)
    print(f"[load] {len(P_t):,} points, {int(mask_wall.sum()):,} wall-face")

    doc = Document()
    set_document_defaults(doc)
    add_header_footer(doc, args.site)

    page_cover(doc, args.site, args.dataset, tilt, plumb, meta)
    page_overview(doc, args.site, args.dataset, tilt, plumb)
    page_plumb(doc, plumb)
    page_tilt(doc, tilt)
    page_horizontal(doc, tilt, P_t, disp, mask_wall, plumb_axis, plumb_origin)
    page_vertical(doc, P_t, disp, mask_wall, plumb_axis, plumb_origin)
    page_conclusion(doc, tilt, P_t, disp, mask_wall)

    out = Path(args.out) if args.out else \
        root / f"report_{datetime.now():%Y%m%d}.docx"
    doc.save(str(out))
    print(f"[done] → {out}")


if __name__ == "__main__":
    main()
